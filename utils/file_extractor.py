from pathlib import Path

import io
import pandas as pd
from PIL import Image

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".csv",
    ".xlsx",
    ".jpg",
    ".jpeg",
    ".png",
}


def extract_text_from_file(uploaded_file) -> str:
    """
    Extract text from a supported uploaded file.

    Supported formats:
        PDF, DOCX, TXT, CSV, XLSX, JPG, JPEG, PNG

    Returns:
        Extracted text as a string.

    Raises:
        ValueError: For unsupported or empty files.
        Exception: For file-specific extraction errors.
    """

    if uploaded_file is None:
        raise ValueError("No file was uploaded.")

    file_name = uploaded_file.name
    extension = Path(file_name).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension or 'unknown'}"
        )

    file_bytes = uploaded_file.getvalue()

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    # --------------------------------------------------
    # PDF
    # --------------------------------------------------
    if extension == ".pdf":
        return _extract_pdf(file_bytes)

    # --------------------------------------------------
    # DOCX
    # --------------------------------------------------
    if extension == ".docx":
        return _extract_docx(file_bytes)

    # --------------------------------------------------
    # TXT
    # --------------------------------------------------
    if extension == ".txt":
        return _extract_txt(file_bytes)

    # --------------------------------------------------
    # CSV
    # --------------------------------------------------
    if extension == ".csv":
        return _extract_csv(file_bytes)

    # --------------------------------------------------
    # XLSX
    # --------------------------------------------------
    if extension == ".xlsx":
        return _extract_xlsx(file_bytes)

    # --------------------------------------------------
    # Images
    # --------------------------------------------------
    if extension in {".jpg", ".jpeg", ".png"}:
        return _extract_image_ocr(file_bytes)

    raise ValueError("Unable to extract text from the uploaded file.")


def _extract_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file."""

    reader = PdfReader(io.BytesIO(file_bytes))

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    extracted_text = "\n".join(pages).strip()

    if not extracted_text:
        raise ValueError(
            "No selectable text could be extracted from the PDF."
        )

    return extracted_text


def _extract_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file."""

    document = Document(io.BytesIO(file_bytes))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    extracted_text = "\n".join(paragraphs).strip()

    if not extracted_text:
        raise ValueError(
            "No text could be extracted from the Word document."
        )

    return extracted_text


def _extract_txt(file_bytes: bytes) -> str:
    """Extract text from a plain text file."""

    try:
        extracted_text = file_bytes.decode("utf-8").strip()
    except UnicodeDecodeError:
        extracted_text = file_bytes.decode(
            "latin-1",
            errors="ignore"
        ).strip()

    if not extracted_text:
        raise ValueError("The text file is empty.")

    return extracted_text


def _extract_csv(file_bytes: bytes) -> str:
    """Extract text from a CSV file."""

    dataframe = pd.read_csv(io.BytesIO(file_bytes))

    if dataframe.empty:
        raise ValueError("The CSV file contains no data.")

    # Convert all cells to strings and combine them
    extracted_text = (
        dataframe.astype(str)
        .fillna("")
        .apply(lambda row: " ".join(row.values), axis=1)
        .str.cat(sep="\n")
        .strip()
    )

    if not extracted_text:
        raise ValueError("No usable text could be extracted from the CSV.")

    return extracted_text


def _extract_xlsx(file_bytes: bytes) -> str:
    """Extract text from all populated cells in an XLSX file."""

    workbook = load_workbook(
        filename=io.BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    lines = []

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            values = [
                str(value).strip()
                for value in row
                if value is not None and str(value).strip()
            ]

            if values:
                lines.append(" ".join(values))

    workbook.close()

    extracted_text = "\n".join(lines).strip()

    if not extracted_text:
        raise ValueError(
            "No usable text could be extracted from the Excel file."
        )

    return extracted_text


def _extract_image_ocr(file_bytes: bytes) -> str:
    """Extract text from JPG/JPEG/PNG using OCR."""

    try:
        import pytesseract
    except ImportError as exc:
        raise ValueError(
            "OCR support is not installed. Install pytesseract first."
        ) from exc

    image = Image.open(io.BytesIO(file_bytes))

    extracted_text = pytesseract.image_to_string(image).strip()

    if not extracted_text:
        raise ValueError(
            "No readable text could be detected in the image."
        )

    return extracted_text